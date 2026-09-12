"""Original three-stage pipeline over persisted source tables; outputs are versioned."""
import io
import json
import uuid
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from data_hub import engine,pack,stamp,objects,table_of

class Run(BaseModel):
    expected_version:str
    confirmed:bool=False

def excel(sheets):
    wb=Workbook();wb.remove(wb.active)
    for name,rows in sheets.items():
        ws=wb.create_sheet(name[:31]);ws.freeze_panes='A2'
        for row in rows: ws.append(row)
        for row in ws:
            for c in row:
                if c.data_type=='f':c.data_type='s'
        for c in ws[1]: c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(horizontal='center')
        for i,c in enumerate(ws[1],1):
            header=str(c.value or '');ws.column_dimensions[c.column_letter].width=min(42,max(15,len(header)*2+3))
            if any(s in header for s in ('率','占比')):
                for row in ws.iter_rows(min_row=2,min_col=i,max_col=i):row[0].number_format='0.00%'
        if name.startswith('利润汇总'):
            for row in ws.iter_rows(min_row=2):
                if len(row)>=3 and any(s in str(row[-2].value) for s in ('率','占比')): row[-1].number_format='0.00%'
    buffer=io.BytesIO();wb.save(buffer);return buffer.getvalue()

def create_router(hub):
    router=APIRouter(prefix='/api/hub/pipeline')
    @router.post('/run')
    def run(payload:Run):
        if not payload.confirmed: raise HTTPException(422,'请确认按当前批次执行原三阶段处理，并激活处理结果')
        current=hub.current()
        if current['version']!=payload.expected_version:raise HTTPException(409,'数据已变化')
        tables=current['data'].get('source_tables',{})
        if any(not tables.get(k) for k in ('orders','settlement','product_pack')):raise HTTPException(422,'缺少马帮订单、TikTok结算或产品包；不以缺少成本的0值生成经营结论')
        orders={};settlements={};packages=tables['product_pack']
        for key,target in [('orders',orders),('settlement',settlements)]:
            for entry in tables[key]:
                shop=entry.get('shop')
                if not shop:raise HTTPException(422,entry['name']+' 缺少店编；请按 MSxxxx 文件名重新导入')
                if shop in target:raise HTTPException(422,'同一店编的同类型原表不止一张，请先合并后导入：'+shop)
                target[shop]=entry
        if set(orders)!=set(settlements):raise HTTPException(422,'马帮订单与结算原表的店编不齐全')
        results=[]
        for shop in sorted(orders):
            opts=[p for p in packages if p.get('shop')==shop] or [p for p in packages if not p.get('shop')]
            if len(opts)!=1:raise HTTPException(422,shop+' 未找到唯一匹配的产品包')
            order,bill,pkg=orders[shop],settlements[shop],opts[0]
            if order['period']!=bill['period']:raise HTTPException(422,shop+' 的订单与结算批次周期不同，请先核实关联范围')
            if pkg['currency']!=bill['currency']:raise HTTPException(422,'产品包成本与结算币种不同，不自动换算')
            out1=engine({'op':'pipeline','stage':'stage1','args':[order,pkg]})
            out2=engine({'op':'pipeline','stage':'stage2','args':[bill,out1,pkg,shop]})
            out3=engine({'op':'pipeline','stage':'stage3','args':[out2,shop]})
            results.append({'shop':shop,'out1':out1,'out2':out2,'out3':out3})
        currencies={v['currency'] for v in settlements.values()};periods={v['period'] for v in settlements.values()}
        if len(currencies)!=1 or len(periods)!=1:raise HTTPException(422,'本次各店必须同周期、同币种，避免混入一个利润口径')
        data=current['data'];data['pipeline_results']=results
        previous_partitions=data.get('partition_store',{}).pop('bill',None)
        data['bill']=engine({'op':'bill_normalize','table':table_of([row for r in results for row in objects(r['out2'])])})
        meta=current['meta'];meta['datasets']['bill']={'period':next(iter(periods)),'currency':next(iter(currencies)),
          'updated_at':stamp(),'origin':'原 SKILL_PIPE.stage1/2/3','row_count':len(data['bill']['rows']),
          'warnings':[r['shop']+'：未匹配产品包 SKU '+str(r['out1']['stats']['unmatched']) for r in results if r['out1']['stats']['unmatched']]}
        meta['last_import']=[]
        if previous_partitions:meta['datasets']['bill']['warnings'].append('流水线重建了当前账单全表；之前账单分区仍保存在旧数据版本，避免把旧分区标为本次结果。')
        with hub.db() as con:
            con.execute('BEGIN IMMEDIATE')
            if con.execute('SELECT version FROM hub_head').fetchone()[0]!=payload.expected_version:raise HTTPException(409,'处理期间数据已更新，请重试')
            version=hub.save(con,data,meta,current['version'])
        return {'version':version,'stores':[r['shop'] for r in results],'warnings':meta['datasets']['bill']['warnings']}

    @router.get('/export/{stage}')
    def export(stage:str,version:str='',shop:str=''):
        current=hub.current(version);results=current['data'].get('pipeline_results',[])
        if stage=='profit' and not results:
            table=current['data']['bill'];i=table['header'].index('店编')
            for store in sorted({str(r[i]) for r in table['rows']}):
                scoped={'header':table['header'],'rows':[r for r in table['rows'] if str(r[i])==store]}
                results.append({'shop':store,'out3':engine({'op':'pipeline','stage':'stage3','args':[scoped,store]})})
        if shop:results=[r for r in results if r['shop']==shop]
        if not results:raise HTTPException(404,'没有处理结果')
        if stage not in ('orders','bill','profit'):raise HTTPException(404)
        sheets={}
        if stage=='profit':
            sheets['利润汇总(全部)']=[['店编','类别','指标','数值']]+[[r['shop'],*row] for r in results for row in r['out3']['summary'][1:]]
            for key,prefix in [('店编维度',False),('产品维度',True),('产品标签',True)]:
                header=results[0]['out3']['sheets'][key][0]
                sheets[key+'(全部)']=[(['店编'] if prefix else [])+header]+[([r['shop']] if prefix else [])+row for r in results for row in r['out3']['sheets'][key][1:-1]]
        else:
            for r in results:
                table=r['out1' if stage=='orders' else 'out2'];sheets[r['shop']]=[table['header'],*table['rows']]
        return Response(excel(sheets),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition':f'attachment; filename="skill_{stage}_{current["version"][:8]}.xlsx"'})
    return router
