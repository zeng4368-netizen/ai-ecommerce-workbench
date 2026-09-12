#!/usr/bin/env python3
"""Export the table data embedded in the local Skill dashboards.

The exporter keeps three fidelity layers:
1. the exact JavaScript/JSON literal copied from the source HTML;
2. a UTF-8 JSON representation preserving values and field order;
3. an Excel workbook for convenient inspection and reuse.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import json
import re
import shutil
import subprocess
from collections import OrderedDict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_FILES = {
    "gmv": "TikTok_GMVMax广告经营管理看板v45_简约版_离线.html",
    "daily": "日销异常1.html",
    "aftersales": "售后数据.html",
    "ads_page": "广告数据.html",
}

AFTERSALES_VARIABLES = [
    "BIZ_RATE",
    "BIZ_DATA",
    "STORE_INFO",
    "DAILY_KPI",
    "STORE_MONTHS",
    "STORE_TABLE",
    "STORE_DETAIL",
    "CAT_TREE",
    "CAT_LEVELS",
    "CAT_LEVEL_NAMES",
    "BILL_DATA",
    "CHECK_MONTHS",
    "CHECKIN",
    "DIM_ROWS",
    "MK_CREATOR_TOTAL",
    "MK_CREATOR_ROWS",
    "MK_CREATOR_DETAIL",
    "MK_CREATOR_MONTHLY",
]

CREATOR_ROW_HEADERS = [
    "达人名称",
    "达人归因GMV",
    "退款金额",
    "预计佣金",
    "视频播放量",
    "视频数",
]

CREATOR_DETAIL_HEADERS = [
    "达人名称",
    "达人归因GMV",
    "达人直播归因GMV",
    "联盟视频归因GMV",
    "退款金额",
    "归因订单数",
    "平均订单金额",
    "联盟商品卡归因GMV",
    "视频数",
    "视频播放量",
    "预计佣金",
]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def unique_output_dir(base: Path) -> Path:
    if not base.exists():
        return base
    idx = 2
    while True:
        candidate = base.with_name(f"{base.name}_{idx}")
        if not candidate.exists():
            return candidate
        idx += 1


def extract_js_assignment(text: str, name: str) -> str | None:
    match = re.search(rf"\b(?:const|let|var)\s+{re.escape(name)}\s*=", text)
    if not match:
        return None

    start = match.end()
    i = start
    depth = 0
    quote: str | None = None
    escaped = False
    line_comment = False
    block_comment = False
    root_started = False
    root_closed = False
    literal_end: int | None = None

    while i < len(text):
        ch = text[i]
        nxt = text[i + 1] if i + 1 < len(text) else ""

        if line_comment:
            if ch in "\r\n":
                line_comment = False
            i += 1
            continue

        if block_comment:
            if ch == "*" and nxt == "/":
                block_comment = False
                i += 2
                if root_closed:
                    literal_end = i
            else:
                i += 1
            continue

        if root_closed:
            if ch.isspace():
                i += 1
                continue
            if ch == "/" and nxt == "*":
                block_comment = True
                i += 2
                continue
            if ch == ";":
                return text[start:i].strip()
            return text[start : literal_end or i].strip()

        if quote:
            if escaped:
                escaped = False
            elif ch == "\\":
                escaped = True
            elif ch == quote:
                quote = None
            i += 1
            continue

        if ch == "/" and nxt == "/":
            line_comment = True
            i += 2
            continue
        if ch == "/" and nxt == "*":
            block_comment = True
            i += 2
            continue
        if ch in "'\"`":
            quote = ch
            i += 1
            continue
        if ch in "([{":
            if depth == 0:
                root_started = True
            depth += 1
        elif ch in ")]}" and depth:
            depth -= 1
            if root_started and depth == 0:
                root_closed = True
                literal_end = i + 1
        elif ch == ";" and depth == 0:
            return text[start:i].strip()
        i += 1

    return None


def strip_outer_comments(expr: str) -> str:
    cleaned = expr.strip()
    while cleaned.startswith("/*"):
        end = cleaned.find("*/")
        if end < 0:
            break
        cleaned = cleaned[end + 2 :].lstrip()
    while cleaned.endswith("*/"):
        start = cleaned.rfind("/*")
        if start < 0:
            break
        cleaned = cleaned[:start].rstrip()
    return cleaned


def parse_js_literal(expr: str) -> Any:
    cleaned = strip_outer_comments(expr)
    try:
        return json.loads(cleaned, object_pairs_hook=OrderedDict)
    except json.JSONDecodeError:
        pass

    node_program = (
        "const fs=require('fs');"
        "const src=fs.readFileSync(0,'utf8');"
        "const value=eval('(' + src + ')');"
        "process.stdout.write(JSON.stringify(value));"
    )
    result = subprocess.run(
        ["node", "-e", node_program],
        input=cleaned,
        text=True,
        encoding="utf-8",
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        raise ValueError(result.stderr.strip() or "Node.js could not parse literal")
    return json.loads(result.stdout, object_pairs_hook=OrderedDict)


def json_dump(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )


def exact_literal_dump(path: Path, expr: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(expr, encoding="utf-8")


def ordered_columns(rows: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row.keys():
            key_text = str(key)
            if key_text not in seen:
                seen.add(key_text)
                columns.append(key_text)
    return columns


def tabularize(value: Any, headers: list[str] | None = None) -> tuple[list[str], list[list[Any]]] | None:
    if isinstance(value, list):
        if not value:
            return (headers or [], [])
        if all(isinstance(item, dict) for item in value):
            dict_rows = value
            columns = ordered_columns(dict_rows)
            return columns, [[row.get(col) for col in columns] for row in dict_rows]
        if all(isinstance(item, (list, tuple)) for item in value):
            width = max((len(item) for item in value), default=0)
            columns = list(headers or [])
            if len(columns) < width:
                columns.extend(f"原始列{i}" for i in range(len(columns) + 1, width + 1))
            return columns, [list(item) for item in value]
        return ["值"], [[item] for item in value]

    if isinstance(value, dict):
        if value and all(not isinstance(item, (dict, list, tuple)) for item in value.values()):
            return ["字段", "值"], [[key, item] for key, item in value.items()]
        rows: list[list[Any]] = []
        for key, item in value.items():
            rows.append([key, item])
        return ["字段", "原始值"], rows

    return ["值"], [[value]]


def excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, int) and abs(value) >= 10**15:
        return str(value)
    return value


def append_sheet(workbook: Workbook, title: str, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    safe_title = re.sub(r"[\\/*?:\[\]]", "_", title)[:31] or "数据"
    original_title = safe_title
    idx = 2
    while safe_title in workbook.sheetnames:
        suffix = f"_{idx}"
        safe_title = f"{original_title[:31-len(suffix)]}{suffix}"
        idx += 1

    worksheet = workbook.create_sheet(safe_title)
    if headers:
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for row in rows:
        worksheet.append([excel_value(value) for value in row])
        for cell in worksheet[worksheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for col_idx in range(1, min(worksheet.max_column, 80) + 1):
        values = [worksheet.cell(row=row_idx, column=col_idx).value for row_idx in range(1, min(worksheet.max_row, 200) + 1)]
        width = min(max((len(str(value)) for value in values if value is not None), default=8) + 2, 42)
        worksheet.column_dimensions[get_column_letter(col_idx)].width = max(width, 10)


def write_workbook(path: Path, sheets: list[tuple[str, list[str], list[list[Any]]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in sheets:
        append_sheet(workbook, title, headers, rows)
    if not workbook.sheetnames:
        workbook.create_sheet("无数据")
    workbook.save(path)


def describe_value(value: Any) -> dict[str, Any]:
    info: dict[str, Any] = {"type": type(value).__name__}
    if isinstance(value, list):
        info["records"] = len(value)
        if value and isinstance(value[0], dict):
            info["fields"] = ordered_columns([row for row in value if isinstance(row, dict)])
        elif value and isinstance(value[0], (list, tuple)):
            info["columns"] = max(len(row) for row in value if isinstance(row, (list, tuple)))
    elif isinstance(value, dict):
        info["keys"] = list(value.keys())
    return info


def find_row_list(payload: Any) -> list[Any] | None:
    if not isinstance(payload, dict):
        return payload if isinstance(payload, list) else None
    for key in ("rows", "data", "records"):
        value = payload.get(key)
        if isinstance(value, list):
            return value
    return None


def export_gmv(source: Path, output: Path, manifest: dict[str, Any]) -> None:
    text = source.read_text(encoding="utf-8")
    section = output / "01_GMV_Max"
    section.mkdir(parents=True, exist_ok=True)

    payload_expr = extract_js_assignment(text, "EMBEDDED_PAYLOAD")
    state_expr = extract_js_assignment(text, "EMBEDDED_STATE")
    template_expr = extract_js_assignment(text, "HEADER_TEMPLATE_XLSX_BASE64")
    if payload_expr is None:
        raise RuntimeError("GMV HTML does not contain EMBEDDED_PAYLOAD")

    payload = parse_js_literal(payload_expr)
    exact_literal_dump(section / "EMBEDDED_PAYLOAD_源代码原文.txt", payload_expr)
    json_dump(section / "GMV_Max_内嵌数据_原值.json", payload)

    rows = find_row_list(payload) or []
    tabular = tabularize(rows)
    if tabular:
        write_workbook(section / "GMV_Max_逐行内嵌数据.xlsx", [("原始数据", tabular[0], tabular[1])])

    if state_expr:
        exact_literal_dump(section / "EMBEDDED_STATE_源代码原文.txt", state_expr)
        json_dump(section / "GMV_Max_看板状态_原值.json", parse_js_literal(state_expr))

    if template_expr:
        template_b64 = parse_js_literal(template_expr)
        (section / "GMV_Max_原始上传表头模板.xlsx").write_bytes(base64.b64decode(template_b64))

    manifest["modules"]["GMV Max"] = {
        "source": str(source),
        "fidelity": "逐行内嵌原始数据 + 原始表头模板",
        "payload": describe_value(payload),
        "rows": len(rows),
    }


def export_daily(source: Path, output: Path, manifest: dict[str, Any]) -> None:
    text = source.read_text(encoding="utf-8")
    match = re.search(r'<script\s+type="application/json"\s+id="dashboardData">(.*?)</script>', text, re.S)
    if not match:
        raise RuntimeError("Daily-sales HTML does not contain dashboardData")

    raw_json_text = match.group(1)
    dashboard = json.loads(raw_json_text, object_pairs_hook=OrderedDict)
    section = output / "02_日销库存"
    section.mkdir(parents=True, exist_ok=True)
    exact_literal_dump(section / "dashboardData_源代码原文.json", raw_json_text)
    json_dump(section / "日销库存_全部内嵌数据_原值.json", dashboard)

    source_raw = dashboard.get("raw") or dashboard.get("RAW") or {}
    source_sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    for key, value in source_raw.items() if isinstance(source_raw, dict) else []:
        table = tabularize(value)
        if table:
            source_sheets.append((str(key), table[0], table[1]))
    write_workbook(section / "日销库存_原始上传表_内嵌数据.xlsx", source_sheets)

    derived_sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    for key, value in dashboard.items():
        if key in {"raw", "RAW"}:
            continue
        table = tabularize(value)
        if table:
            derived_sheets.append((str(key), table[0], table[1]))
    write_workbook(section / "日销库存_看板派生数据.xlsx", derived_sheets)

    manifest["modules"]["日销库存"] = {
        "source": str(source),
        "fidelity": "原始上传表与看板派生对象均保留",
        "dashboard_keys": list(dashboard.keys()),
        "raw_tables": {str(key): describe_value(value) for key, value in source_raw.items()} if isinstance(source_raw, dict) else {},
    }


def export_aftersales_page(label: str, source: Path, output: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    text = source.read_text(encoding="utf-8")
    section = output / label
    literal_dir = section / "源代码原始数据字面量"
    json_dir = section / "解析后的原值JSON"
    literal_dir.mkdir(parents=True, exist_ok=True)
    json_dir.mkdir(parents=True, exist_ok=True)

    parsed: OrderedDict[str, Any] = OrderedDict()
    variable_manifest: OrderedDict[str, Any] = OrderedDict()
    for name in AFTERSALES_VARIABLES:
        expr = extract_js_assignment(text, name)
        if expr is None:
            continue
        exact_literal_dump(literal_dir / f"{name}.txt", expr)
        try:
            value = parse_js_literal(expr)
        except ValueError as exc:
            variable_manifest[name] = {"parse_error": str(exc)}
            continue
        parsed[name] = value
        json_dump(json_dir / f"{name}.json", value)
        variable_manifest[name] = describe_value(value)

    bill = parsed.get("BILL_DATA")
    if isinstance(bill, dict) and isinstance(bill.get("header"), list) and isinstance(bill.get("rows"), list):
        write_workbook(
            section / "利润账单_逐行内嵌数据.xlsx",
            [("账单_已处理", list(bill["header"]), [list(row) for row in bill["rows"]])],
        )

    creator_sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    creator_total = parsed.get("MK_CREATOR_TOTAL")
    if isinstance(creator_total, dict):
        creator_sheets.append(("达人总览", ["字段", "值"], [[key, value] for key, value in creator_total.items()]))
    creator_rows = parsed.get("MK_CREATOR_ROWS")
    if isinstance(creator_rows, list):
        creator_sheets.append(("达人气泡排行原值", CREATOR_ROW_HEADERS, [list(row) for row in creator_rows]))
    creator_detail = parsed.get("MK_CREATOR_DETAIL")
    if isinstance(creator_detail, list):
        creator_sheets.append(("达人关注明细原值", CREATOR_DETAIL_HEADERS, [list(row) for row in creator_detail]))
    if creator_sheets:
        write_workbook(section / "达人数据_逐行内嵌数据.xlsx", creator_sheets)

    summary_sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    for name, value in parsed.items():
        if name in {"BILL_DATA", "MK_CREATOR_TOTAL", "MK_CREATOR_ROWS", "MK_CREATOR_DETAIL", "MK_CREATOR_MONTHLY"}:
            continue
        table = tabularize(value)
        if table:
            summary_sheets.append((name, table[0], table[1]))
    write_workbook(section / "售后看板_内嵌汇总对象.xlsx", summary_sheets)

    combined_json = section / "全部可解析内嵌对象_原值.json"
    json_dump(combined_json, parsed)
    manifest["modules"][label] = {
        "source": str(source),
        "fidelity": "利润账单和达人为逐行数据；售后主看板仅有汇总对象",
        "variables": variable_manifest,
    }
    return parsed


def validate_outputs(output: Path) -> dict[str, Any]:
    validation: dict[str, Any] = {"xlsx": {}, "json": {}, "files": {}}
    for path in sorted(output.rglob("*")):
        if not path.is_file() or path.name == "manifest.json":
            continue
        rel = path.relative_to(output).as_posix()
        validation["files"][rel] = {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
        if path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            validation["xlsx"][rel] = {
                sheet.title: {"rows_with_header": sheet.max_row, "columns": sheet.max_column}
                for sheet in workbook.worksheets
            }
            workbook.close()
        elif path.suffix.lower() == ".json":
            json.loads(path.read_text(encoding="utf-8"))
            validation["json"][rel] = "ok"
    return validation


def semantic_digest(value: Any) -> str:
    return sha256_bytes(json.dumps(value, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def write_readme(output: Path, manifest: dict[str, Any]) -> None:
    modules = manifest["modules"]
    gmv_rows = modules.get("GMV Max", {}).get("rows", 0)
    daily_raw = modules.get("日销库存", {}).get("raw_tables", {})
    after_vars = modules.get("03_售后数据页", {}).get("variables", {})
    bill_rows = after_vars.get("BILL_DATA", {}).get("keys", [])
    creator_records = after_vars.get("MK_CREATOR_DETAIL", {}).get("records", 0)
    daily_lines = "\n".join(
        f"  - `{name}`：{info.get('records', 0)} 行" for name, info in daily_raw.items()
    ) or "  - 未发现原始上传表"

    readme = f"""# Skill 内嵌表格数据原值包

生成日期：{manifest['generated_on']}

这份数据包不是重新计算的结果，而是从当前本地 HTML 看板中直接提取。每个模块尽量同时提供：

1. `源代码原文`：HTML 里的数据字面量原样复制，适合做逐字核对；
2. `原值.json`：字段和值的无损结构化备份；
3. `.xlsx`：便于直接查看、上传或继续处理。

## 当前可交付数据

- GMV Max：{gmv_rows} 行逐行数据，并解码保留了看板自带的原始上传表头模板。
- 日销库存原始上传表：
{daily_lines}
- 售后数据页：利润账单与达人数据可逐行导出；达人关注明细 {creator_records} 行。
- 售后主看板：HTML 中只有 `BIZ_DATA`、`DAILY_KPI`、`STORE_TABLE`、`STORE_DETAIL`、`CAT_TREE` 等汇总对象，没有最初上传的逐单售后明细，因此只能原样交付这些汇总对象。
- “广告数据.html”实际仍是售后/利润旧页面，已独立导出，未与新售后页面强行合并。

## 明确没有内嵌原始表的模块

- `GMV-MAX广告数据清晰` Skill：只有清洗规则、字段规范和脚本，没有内嵌样例业务数据。
- `广告板块数据处理-SKILL说明.md`：只有说明文档，引用的处理脚本和原始业务表不在当前 Skill 文件夹中。
- 售后页面里的“视频数据”和“广告数据”页签：当前是待接入占位模块，没有逐行数据。

## 关于“原封不动”

- JSON 与源代码原文用于保证值和字段顺序可追溯。
- Excel 对超过 15 位的数值会自动丢精度，因此导出时将这类数值按文本写入；其数字字符没有改变。
- 对源代码中只用数组保存、没有同时保存表头的汇总对象，Excel 使用 `原始列1`、`原始列2` 占位；原始数组本身仍完整保存在 JSON 和源代码原文中。
- `manifest.json` 含源文件和每个导出文件的 SHA-256、行列校验结果。
"""
    (output / "README.md").write_text(readme, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Export embedded Skill dashboard tables")
    parser.add_argument("--source", type=Path, default=Path(r"C:\Users\PC\Desktop\Sikll"))
    parser.add_argument("--output-root", type=Path, default=Path(r"D:\codex\data\output"))
    args = parser.parse_args()

    requested = args.output_root / f"skill_embedded_tables_{date.today().isoformat()}"
    output = unique_output_dir(requested)
    output.mkdir(parents=True, exist_ok=False)

    source_paths = {key: args.source / filename for key, filename in SOURCE_FILES.items()}
    missing = [str(path) for path in source_paths.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing source files: " + ", ".join(missing))

    manifest: dict[str, Any] = {
        "generated_on": date.today().isoformat(),
        "source_root": str(args.source),
        "output_root": str(output),
        "source_files": {
            key: {"path": str(path), "bytes": path.stat().st_size, "sha256": sha256_file(path)}
            for key, path in source_paths.items()
        },
        "modules": OrderedDict(),
    }

    export_gmv(source_paths["gmv"], output, manifest)
    export_daily(source_paths["daily"], output, manifest)
    aftersales = export_aftersales_page("03_售后数据页", source_paths["aftersales"], output, manifest)
    ads_page = export_aftersales_page("04_广告数据页", source_paths["ads_page"], output, manifest)

    common = sorted(set(aftersales).intersection(ads_page))
    manifest["cross_page_comparison"] = {
        name: {
            "same_semantic_value": semantic_digest(aftersales[name]) == semantic_digest(ads_page[name]),
            "售后数据页_sha256": semantic_digest(aftersales[name]),
            "广告数据页_sha256": semantic_digest(ads_page[name]),
        }
        for name in common
    }

    write_readme(output, manifest)
    manifest["validation"] = validate_outputs(output)
    json_dump(output / "manifest.json", manifest)

    archive_base = output.parent / output.name
    archive_path = Path(shutil.make_archive(str(archive_base), "zip", root_dir=output.parent, base_dir=output.name))
    print(json.dumps({
        "output": str(output),
        "archive": str(archive_path),
        "archive_bytes": archive_path.stat().st_size,
        "archive_sha256": sha256_file(archive_path),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
