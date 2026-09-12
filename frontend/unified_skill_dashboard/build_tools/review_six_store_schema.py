"""Onboarding review for the explicitly named first six-store batch only.

--inspect reads schemas. --apply validates existing originals only, with an audit
manifest and unchanged export identity/date selectors. Never re-exports or calls AI.
"""
import sys,json,copy,hashlib
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from ziniao_bridge import profile_hash
from ziniao_reports import parse_report,metrics,content_fingerprint
from data_hub import pack,stamp

BATCH='4f568ffe190b4564a3679beac172dd19'

def export_hash(profile):
    p=copy.deepcopy(profile)
    for k in ('reviewed','export_reviewed','validation_note'):p.pop(k,None)
    for r in p.get('reports',{}).values():r.pop('schema',None)
    return profile_hash(p)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    import server
    c=server.collection;b=c.batch(BATCH)
    if c.settings()['enabled']:raise SystemExit('Pause scheduler first')
    if '--inspect' in sys.argv:
        out=[]
        for rid in b['run_ids']:
            j=c.get(rid);p=c.recipe(j['store_id']);f=j['files']['settlement']
            wb=load_workbook(c.hub.data/'raw/ai_workbench/imports'/rid/f['saved_name'],read_only=True)
            name='Order details' if 'Order details' in wb.sheetnames else '订单详情';ws=wb[name];ws.reset_dimensions()
            headers=[str(v or '').strip() for v in next(ws.values)];wb.close()
            out.append({'run_id':rid,'store_id':j['store_id'],'profile_hash':j['profile_hash'],'export_hash':export_hash(p),'file_sha256':f['sha256'],'headers':headers,'sheet':name,'currency':j['currency']})
        print(json.dumps(out,ensure_ascii=False))
    elif '--apply' in sys.argv:
        manifest=json.loads((Path(__file__).resolve().parents[1]/'reports/six_store_onboarding_audit.json').read_text(encoding='utf-8'))
        staged=[]
        for old in manifest:
            j=c.get(old['run_id']);p=c.recipe(j['store_id']);f=j['files']['settlement']
            assert j['batch_id']==BATCH and j['store_id']==old['store_id']
            assert j['profile_hash']==old['profile_hash'] and export_hash(p)==old['export_hash']
            assert j['state'] in ('paused','pending_review')
            content=(c.hub.data/'raw/ai_workbench/imports'/j['id']/f['saved_name']).read_bytes()
            assert hashlib.sha256(content).hexdigest()==old['file_sha256']==f['sha256']
            report=parse_report(content,f['name'],'settlement',p['reports']['settlement']['schema'],j['start'],j['end'],j['currency'])
            j['snapshot']={'store_id':j['store_id'],'shop':j['shop'],'currency':j['currency'],'start':j['start'],'end':j['end'],
                'platform_timezone':p['platform_timezone'],'reports':{'settlement':report},'job_id':j['id'],
                'required_reports':['settlement'],'partial_day':j['end'],'fingerprint':content_fingerprint({'settlement':report})}
            j.setdefault('schema_reviews',[]).append({'at':stamp(),'reason':'首批真实账单表头与平台汇总人工核验；未更改导出身份或日期选择器','old_profile_hash':j['profile_hash'],'new_profile_hash':profile_hash(p),'file_sha256':f['sha256']})
            j.update(profile_hash=profile_hash(p),summary=metrics(j['snapshot']),state='pending_review',error='',updated_at=stamp())
            staged.append(j)
        with c.db() as con:
            con.execute('BEGIN IMMEDIATE')
            for j in staged:
                fresh=json.loads(con.execute('SELECT value FROM ziniao_runs WHERE id=?',(j['id'],)).fetchone()[0])
                if fresh['state'] not in ('paused','pending_review'):raise RuntimeError('Concurrent task change')
                con.execute('UPDATE ziniao_runs SET state=?,value=? WHERE id=?',('pending_review',pack(j),j['id']))
                con.execute('INSERT OR REPLACE INTO hub_jobs VALUES(?,?)',(j['id'],pack({'id':j['id'],'created_at':j['created_at'],'entries':[],'errors':[],'originals':[{'id':k,**v} for k,v in j['files'].items()],'origin':'ziniao','note':'六店首批账单，文件校验与审计已保存'})))
        print(json.dumps([{'shop':j['shop'],'currency':j['currency'],'rows':j['summary']['settlement']['transaction_count'],'settlement':j['summary']['settlement']['settlement']} for j in staged],ensure_ascii=False))
    else:raise SystemExit('Use --inspect or --apply')
