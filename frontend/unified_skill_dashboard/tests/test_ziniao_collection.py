"""Synthetic acceptance data only. Never opens a shop or calls a paid model."""
import copy
import io
import json
import hashlib
from datetime import datetime, timedelta

import pytest
import yaml
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook

from test_server import server
from ziniao_collection import Collection, STORE, MY, bill_period
from ziniao_bridge import Bridge, profile_hash
from ziniao_reports import parse_report, metrics, content_fingerprint, native_query, report_kinds


def spec(kind):
    fields = ({'order_id':'Order ID','line_id':'Line ID','status':'Status','created_at':'Created',
               'quantity':'Quantity','sku':'SKU','currency':'Currency'} if kind=='orders' else
              {'transaction_id':'Transaction ID','order_id':'Order ID','settled_at':'Settled',
               'settlement':'Settlement','refund':'Refund','currency':'Currency'})
    return {'headers':list(fields.values())+['Buyer phone'], 'fields':fields,
            'unique_key':['order_id','line_id'] if kind=='orders' else ['transaction_id'],
            'date_format':'%Y-%m-%d %H:%M:%S','sheet':'Data','header_row':1}


def workbook(kind, change=None, empty=False):
    row = (['1234567890123456789','line-1','Cancelled','2026-09-06 12:00:00',2,'SKU-1','MYR','PRIVATE-PHONE']
        if kind=='orders' else ['TX-1','outside-window','2026-09-07 12:00:00','10.10','-1.20','MYR','PRIVATE-PHONE'])
    if change: change(row)
    wb=Workbook();ws=wb.active;ws.title='Data';ws.append(spec(kind)['headers'])
    if not empty: ws.append(row)
    b=io.BytesIO();wb.save(b);return b.getvalue()


@pytest.fixture
def system(tmp_path,monkeypatch):
    monkeypatch.setattr(server,'DB',tmp_path/'workbench.sqlite3')
    monkeypatch.setattr(server.hub,'data',tmp_path)
    profile={'reviewed':True,'platform_timezone':'Asia/Kuala_Lumpur','reports':{k:{'schema':spec(k)} for k in ('orders','settlement')}}
    path=tmp_path/'selectors.yaml';path.write_text(yaml.safe_dump({'ziniao_collection':profile}),encoding='utf-8')
    calls=[]
    def model(messages):
        calls.append(messages)
        assert 'PRIVATE-PHONE' not in json.dumps(messages)
        return {'message':{'content':'先核对窗口外交易。不能将结算金额视为利润。'},'model':'test',
                'provider':'test','usage':{'prompt_tokens':300,'completion_tokens':40},'finish_reason':'stop'}
    c=Collection(server.hub,path,model)
    def collect(job,save,raw_dir,shots):
        raw_dir.mkdir(parents=True,exist_ok=True)
        for k in report_kinds(job):
            data=workbook(k);(raw_dir/(k+'.xlsx')).write_bytes(data)
            job['files'][k]={'name':k+'.xlsx','saved_name':k+'.xlsx','sha256':hashlib.sha256(data).hexdigest(),'size':len(data)}
            save(job)
        return profile
    monkeypatch.setattr(c.bridge,'collect',collect)
    return c,calls,profile


def ready(c):
    # Frozen legacy pilot: retain two-report regression coverage after the policy change.
    job=c.new_run(now=datetime(2026,9,11,9,tzinfo=MY))
    job.update(required_reports=['orders','settlement'],start='2026-09-04',end='2026-09-10',partial_day=None)
    c.save(job);c.execute(job['id'])
    assert c.get(job['id'])['state']=='pending_review'
    return c.get(job['id'])


def test_native_parse_precision_privacy_dates_currency():
    def parse(content):return parse_report(content,'native.xlsx','orders',spec('orders'),'2026-09-04','2026-09-10','MYR')
    r=parse(workbook('orders'))
    assert r['rows'][0]['order_id']=='1234567890123456789'
    assert 'PRIVATE-PHONE' not in json.dumps(r)
    for content in [workbook('orders',lambda row:row.__setitem__(0,1234567890123456789)),
                    workbook('orders',lambda row:row.__setitem__(3,'2026-08-01 00:00:00')),
                    workbook('orders',lambda row:row.__setitem__(6,'THB')),
                    workbook('orders',lambda row:row.__setitem__(4,'NaN')),
                    workbook('orders',empty=True)]:
        with pytest.raises(HTTPException):parse(content)


def test_live_patrol_lease_blocks_bill_export(system):
    import time
    from tiktok_patrol_live import LEASE_SCHEMA
    c,calls,_=system
    job=c.new_run(now=datetime(2026,9,11,9,tzinfo=MY))
    with c.db() as con:
        con.execute(LEASE_SCHEMA)
        con.execute('INSERT INTO ziniao_browser_lease VALUES(1,?,?)',('synthetic-patrol',time.time()+180))
    with pytest.raises(HTTPException) as error:c.submit(job['id'])
    assert error.value.status_code==409
    assert c.get(job['id'])['state']=='queued' and not calls


def test_atomic_import_report_chat_action_and_original_unchanged(system):
    c,calls,_=system;before=c.hub.current();job=ready(c)
    assert c.hub.current()['version']==before['version']
    result=c.activate(job['id'],before['version'])
    assert result['state']=='complete' and result['conversation_id'] and len(calls)==1
    now=c.hub.current()
    for k in before['data']:assert before['data'][k]==now['data'][k]
    assert metrics(now['data']['tiktok_native'][STORE])['settlement']['unmatched_transaction_count']==1
    assert metrics(now['data']['tiktok_native'][STORE])['settlement']['revenue'] is None
    rows=native_query(now,STORE,'orders')['rows'];assert rows[0]['source_row']==2
    api=TestClient(server.app)
    chat=api.get('/api/hub/conversations/'+result['conversation_id']).json()
    assert chat['pinned_version']==now['version']
    turn=chat['turns'][0]
    response=api.post('/api/actions/from-analysis',json={'analysis_id':turn['analysis_id'],'conversation_id':chat['id'],
        'module':'finance','entity':'EXPOSE TK','title':'核对结算窗口','action':'核对窗口外订单','confirmed':True})
    assert response.status_code==200,response.text
    assert response.json()['item']['state']=='candidate'
    assert response.json()['item']['data_version']==now['version']


def test_duplicate_revision_conflict_and_quota(system):
    c,calls,_=system;job=ready(c);c.activate(job['id'],job['expected_version']);version=c.hub.current()['version']
    second=ready(c);result=c.activate(second['id'],version)
    assert result['duplicate'] and result['version']==version and len(calls)==1
    third=ready(c)
    with pytest.raises(HTTPException):c.activate(third['id'],'missing-version')
    assert c.get(third['id'])['state']=='pending_review'
    third['snapshot']['reports']['settlement']['rows'][0]['settlement']='12.00'
    third['snapshot']['fingerprint']=content_fingerprint(third['snapshot']['reports']);c.save(third)
    result=c.activate(third['id'],version)
    assert result['version']!=version and result['ai_status']=='daily_limit' and len(calls)==1
    assert c.hub.current(version)['data']['tiktok_native'][STORE]['reports']['settlement']['rows'][0]['settlement']=='10.10'


def test_partial_failure_does_not_activate(system,monkeypatch):
    c,_,_=system;before=c.hub.current()['version']
    def broken(*args):raise HTTPException(409,'验证码')
    monkeypatch.setattr(c.bridge,'collect',broken)
    j=c.new_run();c.execute(j['id'])
    assert c.get(j['id'])['state']=='paused' and c.hub.current()['version']==before
    with pytest.raises(HTTPException):c.new_run()


def test_ai_timeout_no_retry_data_stays_valid(system):
    c,_,_=system
    def timeout(messages):raise TimeoutError()
    c.model=timeout;j=ready(c);r=c.activate(j['id'],j['expected_version'])
    assert r['ai_status']=='failed' and r['state']=='complete'
    assert c.analyze(j['id'])['ai_status']=='daily_limit'


def test_scheduler_no_offline_catchup_and_changed_recipe(system):
    c,_,profile=system
    due=datetime(2026,9,11,9,tzinfo=MY)
    c.setting(enabled=True,next_due=due.isoformat(),approved_profile_hash=profile_hash(profile))
    c.tick(due+timedelta(hours=1))
    assert c.settings()['last_schedule']['status']=='missed' and c.status()['runs']==[]
    c.setting(next_due=due.isoformat(),approved_profile_hash='outdated')
    c.tick(due)
    assert c.settings()['enabled'] is False


def test_doctor_warning_fails_closed(system,monkeypatch):
    c,_,_=system
    out='✓ 配置文件: local\n✓ API Key 有效\n✓ 客户端登录用户: demo\n✓ ZClaw Bridge 连通正常\n⚠ 当前客户端版本不支持终端绑定检查\n✓ 全部检查通过'
    monkeypatch.setattr(c.bridge,'run',lambda *a,**kw:(0,out))
    assert not c.check()['ready']
    assert 'API Key' not in json.dumps(c.status()['plan']['last_check'].get('message'))


def test_native_empty_optional_field_and_duplicate_key():
    schema=spec('orders');schema['empty_export_verified']=True
    r=parse_report(workbook('orders',empty=True),'empty.xlsx','orders',schema,'2026-09-04','2026-09-10','MYR')
    assert r['row_count']==0
    text=','.join(schema['headers'])+'\n'+'o,l,Done,2026-09-05 00:00:00,0,s,MYR,p\n'*2
    with pytest.raises(HTTPException):parse_report(text.encode(),'repeat.csv','orders',schema,'2026-09-04','2026-09-10','MYR')


def test_piped_cli_success_on_stderr(system,monkeypatch):
    from types import SimpleNamespace
    import ziniao_bridge
    c,_,_=system
    monkeypatch.setattr(c.bridge,'command',lambda:['node','cli.js'])
    monkeypatch.setattr(ziniao_bridge.subprocess,'run',lambda *a,**k:SimpleNamespace(returncode=0,stdout=b'',stderr='✓ 页面已导航\n'.encode()))
    assert c.bridge.run(['page','visit','--store-id',STORE,'--url','https://example.invalid'])['confirmed']


def test_stale_excel_dimensions_do_not_drop_rows_or_columns():
    from zipfile import ZipFile,ZIP_DEFLATED
    import re
    b=Workbook();s=b.active;s.title='Data';s.append(spec('orders')['headers'])
    s.append(['long-id-1','l1','Cancelled','2026-09-06 12:00:00',0,'S1','MYR','PRIVATE'])
    s.append(['long-id-2','l2','Done','2026-09-07 12:00:00',3,'S2','MYR','PRIVATE'])
    raw=io.BytesIO();b.save(raw);out=io.BytesIO()
    with ZipFile(raw) as src,ZipFile(out,'w',ZIP_DEFLATED) as dst:
        for info in src.infolist():
            value=src.read(info.filename)
            if info.filename=='xl/worksheets/sheet1.xml':value=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:B2"',value)
            dst.writestr(info,value)
    r=parse_report(out.getvalue(),'data.xlsx','orders',spec('orders'),'2026-09-04','2026-09-10','MYR')
    assert r['row_count']==2 and r['rows'][1]['quantity']=='3'


def test_settlement_control_sheet_and_blank_separator():
    schema=spec('settlement');schema['headers'].append('');schema['control_sheet']='Report'
    schema['control_totals']={'settlement':'结算总金额'}
    b=Workbook();s=b.active;s.title='Data';s.append(schema['headers'])
    s.append(['tx1','o1','2026-09-07 12:00:00','10.10','-1.20','MYR','PRIVATE',''])
    ws=b.create_sheet('Report')
    for row in [['时间范围','2026/09/04-2026/09/10'],['时区','UTC+8'],['货币','MYR'],['结算总金额','10.10']]:ws.append(row)
    out=io.BytesIO();b.save(out)
    assert parse_report(out.getvalue(),'data.xlsx','settlement',schema,'2026-09-04','2026-09-10','MYR')['control']['结算总金额']=='10.10'
    ws['B4']='999';out=io.BytesIO();b.save(out)
    with pytest.raises(HTTPException):parse_report(out.getvalue(),'data.xlsx','settlement',schema,'2026-09-04','2026-09-10','MYR')


def test_ai_parallel_guard_and_interrupted_request(system):
    import time
    from ziniao_collection import local_now
    c,calls,_=system;j=ready(c)
    c.model=lambda m: (_ for _ in ()).throw(TimeoutError())
    c.activate(j['id'],j['expected_version']);j=c.get(j['id'])
    j.update(ai_status='running',ai_started_at=time.time());c.save(j)
    with c.db() as con:con.execute("UPDATE ziniao_ai_calls SET status='started'")
    with pytest.raises(HTTPException):c.analyze(j['id'],manual=True)
    assert c.get(j['id'])['ai_status']=='running'
    j['ai_started_at']=time.time()-601;c.save(j)
    c.analyze(j['id'])
    assert c.get(j['id'])['ai_status']=='running'  # No paid automatic retry after a restart.
    assert c.analyze(j['id'],manual=True)['ai_status']=='failed'


def test_order_url_cross_month_and_receipt_scope():
    from ziniao_tiktok import order_url,receipt_time
    from urllib.parse import parse_qs,urlsplit
    j={'start':'2026-08-27','end':'2026-09-02'}
    q=parse_qs(urlsplit(order_url({'url':'https://seller-my.tiktok.com/order'},j)).query)
    a,z=map(int,q['time_order_created[]']);assert z-a+1==7*86400*1000
    assert datetime.fromtimestamp(a/1000,MY).strftime('%Y-%m-%d %H:%M')=='2026-08-27 00:00'
    assert receipt_time('old.xlsx','orders') is None
    assert receipt_time('income_20260911122255(UTC+8).xlsx','settlement') is not None


def test_full_product_risks_and_reordered_fingerprint(system):
    from ziniao_reports import operating_breakdown
    c,_,_=system;j=ready(c);snap=j['snapshot']
    snap['reports']['orders']['rows']=[{**snap['reports']['orders']['rows'][0],'order_id':str(i),'source_row':i+2} for i in range(10)]
    risk=operating_breakdown(snap)['anomalies'];assert len(risk)==1 and risk[0]['cancelled_lines']==10
    before=content_fingerprint(snap['reports'])
    snap['reports']['orders']['rows'].reverse()
    assert content_fingerprint(snap['reports'])==before


@pytest.mark.parametrize('iso,expected',[
    ('2026-09-11T09:00:00+08:00',('2026-09-01','2026-09-11')),
    ('2026-09-23T09:00:00+08:00',('2026-09-01','2026-09-23')),
    ('2026-10-21T09:00:00+08:00',('2026-10-01','2026-10-21')),
    ('2026-10-01T00:00:00+08:00',('2026-10-01','2026-10-01')),
    ('2028-02-29T09:00:00+08:00',('2028-02-01','2028-02-29')),
    ('2026-12-31T16:01:00+00:00',('2027-01-01','2027-01-01')),
    ('2026-09-11T09:00:00',('2026-09-01','2026-09-11')),
])
def test_bill_period(iso,expected):
    assert bill_period(datetime.fromisoformat(iso))==expected


@pytest.mark.parametrize('origin',['manual','scheduled'])
def test_new_bill_only_pipeline(system,origin):
    c,calls,_=system;before=c.hub.current()
    j=c.new_run(origin=origin,now=datetime(2026,9,11,9,tzinfo=MY))
    assert (j['start'],j['end'],j['partial_day'])==('2026-09-01','2026-09-11','2026-09-11')
    assert report_kinds(j)==['settlement'] and j['report_name']=='账单'
    c.execute(j['id']);j=c.get(j['id'])
    assert j['state']=='pending_review' and set(j['files'])=={'settlement'}
    assert j['summary']['orders']['order_count'] is None
    assert j['summary']['settlement']['unmatched_transaction_count'] is None
    assert j['summary']['daily']==[] and len(calls)==0
    c.activate(j['id'],j['expected_version'])
    current=c.hub.current();snap=current['data']['tiktok_native'][STORE]
    assert set(snap['reports'])=={'settlement'}
    assert native_query(current,STORE,'settlement')['rows'][0]['settlement']=='10.10'
    for kind in ('orders','products','anomalies'):
        result=native_query(current,STORE,kind)
        assert result['available'] is False and result['total'] is None
    for k in before['data']:assert before['data'][k]==current['data'][k]
    assert len(calls)==1 and 'orders_available' in json.dumps(calls[0])
    # Reopening the service does not recalculate a frozen job's dates.
    reopened=Collection(c.hub,c.bridge.selectors,c.model)
    assert reopened.get(j['id'])['end']=='2026-09-11'


def test_bill_includes_today_and_rejects_outside_month():
    def parse(day):
        return parse_report(workbook('settlement',lambda r:r.__setitem__(2,day+' 12:00:00')),
            'bill.xlsx','settlement',spec('settlement'),'2026-09-01','2026-09-11','MYR')
    assert parse('2026-09-01')['row_count']==1
    assert parse('2026-09-11')['row_count']==1
    for day in ('2026-08-31','2026-09-12'):
        with pytest.raises(HTTPException):parse(day)


def test_policy_migration_preserves_legacy_jobs_and_schedule(system):
    c,_,_=system;j=ready(c);j.pop('required_reports');c.save(j)
    before=c.get(j['id']);settings=c.settings()
    settings.pop('collection_policy');settings.update(days=7,approved_profile_hash='approved',enabled=False)
    with c.db() as con:con.execute('UPDATE ziniao_settings SET value=? WHERE id=1',(json.dumps(settings),))
    migrated=c.settings()
    assert migrated['required_reports']==['settlement'] and migrated['days'] is None
    assert migrated['approved_profile_hash']=='approved' and not migrated['enabled']
    assert c.get(j['id'])==before and report_kinds(before)==['orders','settlement']


def test_collector_only_enters_bill_recipe(system,monkeypatch,tmp_path):
    import ziniao_tiktok
    c,_,profile=system;j=c.new_run(now=datetime(2026,9,11,9,tzinfo=MY));seen=[]
    monkeypatch.setattr(c.bridge,'check',lambda:{'ready':True})
    def checked(kinds,store_id):
        assert kinds==['settlement'];return profile
    monkeypatch.setattr(c.bridge,'profile',checked)
    monkeypatch.setattr(c.bridge,'run',lambda args:({'name':j['shop'],'storeId':STORE} if args[0]=='store' else
        {'running':True,'downloadFolderPath':str(tmp_path)}))
    class UI:
        def __init__(self,*args):pass
        def prepare(self,kind,resume):
            seen.append(kind);raise HTTPException(409,'TEST: stop before any export')
    monkeypatch.setattr(ziniao_tiktok,'TikTokUI',UI)
    with pytest.raises(HTTPException):ziniao_tiktok.collect(c.bridge,j,c.save,tmp_path,tmp_path)
    assert seen==['settlement']
