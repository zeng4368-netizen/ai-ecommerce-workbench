#!/usr/bin/env python3
"""
GMV MAX 广告数据清洗汇总脚本

功能：
  1. 读取多家店铺的 TikTok GMV MAX 广告数据文件（Excel/CSV），文件名以店编命名
  2. 读取店铺人员对应名单（店名|店编|国家|初级|中级|储高/见高|CEO）
  3. 对每家店铺执行数据清洗与汇总：
     - 剔除 Cost（广告消耗）为 0 的行
     - 统计总素材数、总消耗、平均 CTR / CVR
     - 按 Time posted 筛选近 7 天（L7D）新建素材，统计数量、消耗、CTR / CVR
  4. 合并人员名单信息，输出汇总 Excel 表格

用法：
  python gmv_max_cleaner.py --ad-dir <广告数据目录> --mapping <人员名单> [--start-date 2024-07-01] [--output 汇总结果.xlsx]

  或指定具体文件：
  python gmv_max_cleaner.py --ad-files <文件1> <文件2> ... --mapping <人员名单> [--start-date 2024-07-01] [--output 汇总结果.xlsx]
"""

import argparse
import os
import re
import sys
from pathlib import Path

import pandas as pd
import numpy as np

# ============================================================
# 列名别名表 —— 支持 TikTok 导出的中英文列名变体
# ============================================================

AD_COL_ALIASES = {
    'cost': [
        'Cost', 'Spend', '广告消耗', '消耗', 'Ad cost', 'ad_cost',
        'Cost (advertising)', 'Advertising cost', '广告花费',
    ],
    'creative_type': [
        'Creative type', 'Creative Type', '素材类型', 'creative_type',
        'Ad type', 'Ad Type', '广告类型', 'Type',
    ],
    'time_posted': [
        'Time posted', 'Time Posted', '发布时间', '创建时间',
        'Creation time', 'time_posted', 'Created time', 'Creation date',
        'Posted time', 'Posting time',
    ],
    'ctr': [
        'Product ad click rate', 'CTR', '点击率',
        'Click-through rate', 'Click through rate', 'ctr',
        'Product ad CTR', 'Product click rate', 'Click rate',
        '广告点击率', '商品点击率',
    ],
    'cvr': [
        'Ad conversion rate', 'CVR', '转化率',
        'Conversion rate', 'conversion rate', 'cvr',
        'Ad CVR', '广告转化率',
    ],
}

MAP_COL_ALIASES = {
    'store_name': ['店名', 'Store Name', '店铺名称', 'store_name', 'Store name'],
    'store_code': ['店编', 'Store ID', '店铺编号', 'Store Code', 'store_code', 'store_id', '店铺ID'],
    'country': ['国家', 'Country', 'Country/Region', 'country', '国家/地区'],
    'junior': ['初级', 'Junior', 'junior'],
    'mid': ['中级', 'Mid', 'mid', 'Intermediate', 'intermediate'],
    'senior': ['储高/见高', '储高', '见高', 'Senior', 'senior', '储备干部'],
    'ceo': ['CEO', 'ceo', '负责人', 'Owner', 'owner', '管理者'],
}

OUTPUT_COLUMNS = [
    '店名', '店编', '国家', '初级', '中级', '储高/见高', 'CEO',
    'L7D新建素材数', 'L7D新建素材消耗额', '总素材数', '总消耗',
    'L7D CTR', 'L7D CVR', '总CTR', '总CVR',
]


# ============================================================
# 工具函数
# ============================================================

def find_column(columns, aliases):
    """在 DataFrame 列名中查找与别名匹配的列。"""
    col_list = list(columns)
    # 精确匹配
    for alias in aliases:
        for col in col_list:
            if str(col).strip() == alias:
                return col
    # 大小写不敏感匹配
    for alias in aliases:
        for col in col_list:
            if str(col).strip().lower() == alias.lower():
                return col
    # 包含匹配（别名是列名的子串）
    for alias in aliases:
        for col in col_list:
            if alias.lower() in str(col).strip().lower():
                return col
    return None


def parse_percentage(value):
    """解析比率值，统一返回「小数形式」(fraction)。

    支持三种输入并自动归一化：
      - '5.23%' 字符串   -> 0.0523（带百分号，必定视为百分比）
      - 12.34 (浮点数，>1，视为已是百分比数值) -> 0.1234
      - 0.1234 (浮点数，<=1，视为小数)        -> 0.1234
    """
    if pd.isna(value):
        return np.nan
    if isinstance(value, (int, float)):
        raw = float(value)
        # 纯数字：>1 视为百分比数值，<=1 视为小数
        return raw / 100.0 if raw > 1 else raw
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return np.nan
        had_pct = s.endswith('%')
        s = s.rstrip('%').strip()
        try:
            raw = float(s)
        except ValueError:
            return np.nan
        # 带百分号或数值>1 -> 视为百分比；否则视为小数
        if had_pct or raw > 1:
            return raw / 100.0
        return raw
    return np.nan


def fmt_pct(mean_val):
    """将小数均值转换为「百分比数值」(百分点，如 3.25)，保留两位小数。

    返回的是数值(3.25)而非文本；配合 Excel 数字格式 '0.00"%"' 即可显示为 '3.25%'，
    同时单元格仍是真实数字，可在 Excel 中排序、制图、二次计算。
    """
    if mean_val is None or (isinstance(mean_val, float) and pd.isna(mean_val)):
        mean_val = 0.0
    return round(mean_val * 100, 2)


def parse_numeric(value):
    """解析数值，清除货币符号、千位分隔符等。"""
    if pd.isna(value):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return 0.0
        # 去除货币符号、逗号、空格
        value = re.sub(r'[^\d.\-]', '', value)
        try:
            return float(value) if value else 0.0
        except ValueError:
            return 0.0
    return 0.0


def parse_date_column(series):
    """解析日期时间列，支持多种常见格式。"""
    # 先用 pandas 自动解析
    parsed = pd.to_datetime(series, errors='coerce')
    if parsed.notna().sum() > 0:
        return parsed
    # 尝试常见格式
    for fmt in [
        '%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%Y/%m/%d',
        '%d/%m/%Y', '%m/%d/%Y', '%Y年%m月%d日',
        '%Y.%m.%d', '%d.%m.%Y', '%Y%m%d',
    ]:
        try:
            parsed = pd.to_datetime(series, format=fmt, errors='coerce')
            if parsed.notna().sum() > 0:
                return parsed
        except (ValueError, TypeError):
            continue
    return parsed


def read_data_file(filepath):
    """读取 Excel 或 CSV 文件，返回 DataFrame。自动处理 TikTok 导出的表头偏移。"""
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"文件不存在: {filepath}")

    if filepath.suffix.lower() == '.csv':
        # 尝试不同编码
        for encoding in ['utf-8-sig', 'utf-8', 'gbk', 'gb2312', 'latin-1']:
            try:
                df = pd.read_csv(filepath, encoding=encoding)
                if len(df) > 0:
                    return _try_fix_header(df, filepath)
            except (UnicodeDecodeError, pd.errors.EmptyDataError):
                continue
        raise ValueError(f"无法读取 CSV 文件: {filepath}")

    elif filepath.suffix.lower() in ('.xlsx', '.xls'):
        # 先尝试直接读取第一个 sheet
        try:
            df = pd.read_excel(filepath, sheet_name=0, engine='openpyxl')
            return _try_fix_header(df, filepath)
        except Exception:
            pass
        # 尝试跳过前几行（TikTok 导出可能有元数据行）
        for skip in range(1, 6):
            try:
                df = pd.read_excel(filepath, sheet_name=0, skiprows=skip, engine='openpyxl')
                if len(df.columns) > 3 and len(df) > 0:
                    return _try_fix_header(df, filepath)
            except Exception:
                continue
        raise ValueError(f"无法读取 Excel 文件: {filepath}")

    else:
        raise ValueError(f"不支持的文件格式: {filepath.suffix}")


def _try_fix_header(df, filepath):
    """检查并修正表头行（处理 TikTok 导出前几行可能是元数据的情况）。"""
    # 检查当前列是否包含已知列名
    all_aliases = []
    for aliases in AD_COL_ALIASES.values():
        all_aliases.extend(aliases)

    current_score = sum(
        1 for col in df.columns
        if any(str(col).strip().lower() == a.lower() for a in all_aliases)
    )

    if current_score >= 2:
        # 当前表头看起来正确
        return _remove_total_rows(df)

    # 尝试将前几行作为表头
    for i in range(min(5, len(df))):
        candidate_header = df.iloc[i].tolist()
        score = sum(
            1 for val in candidate_header
            if pd.notna(val) and any(
                str(val).strip().lower() == a.lower() for a in all_aliases
            )
        )
        if score >= 2:
            new_df = df.iloc[i + 1:].copy()
            new_df.columns = candidate_header
            new_df.reset_index(drop=True, inplace=True)
            return _remove_total_rows(new_df)

    return _remove_total_rows(df)


def _remove_total_rows(df):
    """移除可能的合计/汇总行。"""
    if len(df) == 0:
        return df
    # 检查最后一行是否是合计行（常见标记：Total / 总计 / 合计）
    last_row = df.iloc[-1]
    first_val = str(last_row.iloc[0]).strip().lower() if pd.notna(last_row.iloc[0]) else ''
    total_keywords = ['total', '总计', '合计', '汇总', 'sum']
    if any(kw in first_val for kw in total_keywords):
        df = df.iloc[:-1].copy()
    # 移除全空行
    df = df.dropna(how='all').copy()
    df.reset_index(drop=True, inplace=True)
    return df


def read_mapping_file(filepath):
    """读取人员名单文件，返回标准化的 DataFrame。"""
    df = read_data_file(filepath)

    # 查找映射列
    col_map = {}
    for key, aliases in MAP_COL_ALIASES.items():
        col = find_column(df.columns, aliases)
        if col is not None:
            col_map[key] = col

    # 检查必要列
    required = ['store_code']
    missing = [k for k in required if k not in col_map]
    if missing:
        raise ValueError(
            f"人员名单缺少必要列: {missing}。"
            f"请确认表头包含: 店名|店编|国家|初级|中级|储高/见高|CEO"
        )

    # 构建标准化 DataFrame
    result = pd.DataFrame()
    for key in ['store_name', 'store_code', 'country', 'junior', 'mid', 'senior', 'ceo']:
        if key in col_map:
            result[key] = df[col_map[key]].astype(str).str.strip()
        else:
            result[key] = ''

    # 清理 NaN
    result = result.fillna('')
    return result


def extract_store_code(filename, mapping_codes):
    """从文件名中提取店编。优先匹配人员名单中的店编。"""
    basename = Path(filename).stem

    # 优先：在文件名中查找匹配名单的店编
    for code in mapping_codes:
        code_str = str(code).strip()
        if code_str and code_str in basename:
            return code_str

    # 备选：提取最长的数字串
    numbers = re.findall(r'\d+', basename)
    if numbers:
        return max(numbers, key=len)

    # 最终备选：使用文件名
    return basename


def clean_and_summarize(ad_df, mapping_row, start_date=None):
    """
    对单家店铺的广告数据执行清洗与汇总。

    参数:
        ad_df: 广告数据 DataFrame
        mapping_row: 该店铺的人员名单信息 (dict)
        start_date: L7D 起始日期 (pd.Timestamp 或 None)

    返回:
        包含所有输出列的 dict
    """
    result = {col: '' for col in OUTPUT_COLUMNS}
    # 填入人员信息
    result['店名'] = mapping_row.get('store_name', '')
    result['店编'] = mapping_row.get('store_code', '')
    result['国家'] = mapping_row.get('country', '')
    result['初级'] = mapping_row.get('junior', '')
    result['中级'] = mapping_row.get('mid', '')
    result['储高/见高'] = mapping_row.get('senior', '')
    result['CEO'] = mapping_row.get('ceo', '')

    # 默认值
    for key in ['L7D新建素材数', 'L7D新建素材消耗额', '总素材数', '总消耗']:
        result[key] = 0
    for key in ['L7D CTR', 'L7D CVR', '总CTR', '总CVR']:
        result[key] = '0.00%'

    # 查找列
    cost_col = find_column(ad_df.columns, AD_COL_ALIASES['cost'])
    creative_type_col = find_column(ad_df.columns, AD_COL_ALIASES['creative_type'])
    time_posted_col = find_column(ad_df.columns, AD_COL_ALIASES['time_posted'])
    ctr_col = find_column(ad_df.columns, AD_COL_ALIASES['ctr'])
    cvr_col = find_column(ad_df.columns, AD_COL_ALIASES['cvr'])

    # 检查必要列
    if not cost_col:
        print(f"  [错误] 未找到 Cost/消耗 列，跳过此店铺")
        return result

    df = ad_df.copy()

    # ===== 步骤 1: 剔除 Cost = 0 的行 =====
    df['_cost'] = df[cost_col].apply(parse_numeric)
    total_before = len(df)
    df = df[df['_cost'] > 0].copy()
    total_after = len(df)
    print(f"  步骤1 - 剔除Cost=0: {total_before} → {total_after} 条 (剔除 {total_before - total_after} 条)")

    if total_after == 0:
        print(f"  [警告] 剔除Cost=0后无数据")
        return result

    # ===== 步骤 2: 统计 Creative type 分布 =====
    if creative_type_col:
        type_counts = df[creative_type_col].fillna('未知').value_counts()
        print(f"  步骤2 - 素材类型分布:")
        for ct, count in type_counts.items():
            print(f"    {ct}: {count} 条")

    # ===== 步骤 3: 按 Time posted 筛选 L7D =====
    if time_posted_col:
        df['_time'] = parse_date_column(df[time_posted_col])
    else:
        df['_time'] = pd.NaT

    if start_date is None:
        # 自动检测: 用数据中最大日期 - 6 天作为 L7D 起点
        valid_times = df['_time'].dropna()
        if len(valid_times) > 0:
            max_date = valid_times.max()
            start_date = (max_date - pd.Timedelta(days=6)).normalize()
            print(f"  步骤3 - 自动检测L7D起始日期: {start_date.strftime('%Y-%m-%d')} (数据最大日期: {max_date.strftime('%Y-%m-%d')})")
        else:
            print(f"  步骤3 - [警告] 无法解析发布时间，L7D指标将为0")
            start_date = None
    else:
        if isinstance(start_date, str):
            start_date = pd.to_datetime(start_date)
        start_date = pd.Timestamp(start_date).normalize()
        print(f"  步骤3 - 使用指定L7D起始日期: {start_date.strftime('%Y-%m-%d')}")

    # 筛选 L7D 数据
    if start_date is not None:
        l7d_mask = df['_time'] >= start_date
        l7d_df = df[l7d_mask].copy()
    else:
        l7d_df = pd.DataFrame()

    # ===== 步骤 4: 计算 CTR / CVR =====
    # 总体 CTR / CVR
    if ctr_col:
        total_ctr_series = df[ctr_col].apply(parse_percentage)
        total_ctr = total_ctr_series.dropna().mean() if total_ctr_series.notna().any() else 0
    else:
        print(f"  步骤4 - [警告] 未找到CTR列 (Product ad click rate)")
        total_ctr = 0

    if cvr_col:
        total_cvr_series = df[cvr_col].apply(parse_percentage)
        total_cvr = total_cvr_series.dropna().mean() if total_cvr_series.notna().any() else 0
    else:
        print(f"  步骤4 - [警告] 未找到CVR列 (Ad conversion rate)")
        total_cvr = 0

    # L7D CTR / CVR
    if len(l7d_df) > 0 and ctr_col:
        l7d_ctr_series = l7d_df[ctr_col].apply(parse_percentage)
        l7d_ctr = l7d_ctr_series.dropna().mean() if l7d_ctr_series.notna().any() else 0
    else:
        l7d_ctr = 0

    if len(l7d_df) > 0 and cvr_col:
        l7d_cvr_series = l7d_df[cvr_col].apply(parse_percentage)
        l7d_cvr = l7d_cvr_series.dropna().mean() if l7d_cvr_series.notna().any() else 0
    else:
        l7d_cvr = 0

    # ===== 汇总结果 =====
    result['总素材数'] = total_after
    result['总消耗'] = round(df['_cost'].sum(), 2)
    result['L7D新建素材数'] = len(l7d_df)
    result['L7D新建素材消耗额'] = round(l7d_df['_cost'].sum(), 2) if len(l7d_df) > 0 else 0
    result['L7D CTR'] = fmt_pct(l7d_ctr)
    result['L7D CVR'] = fmt_pct(l7d_cvr)
    result['总CTR'] = fmt_pct(total_ctr)
    result['总CVR'] = fmt_pct(total_cvr)

    print(f"  汇总 → 总素材: {result['总素材数']} 条 | 总消耗: {result['总消耗']} | "
          f"L7D素材: {result['L7D新建素材数']} 条 | L7D消耗: {result['L7D新建素材消耗额']} | "
          f"总CTR: {result['总CTR']}% | 总CVR: {result['总CVR']}%")

    return result


def write_output(results, output_path):
    """将汇总结果写入 Excel 文件。"""
    df = pd.DataFrame(results, columns=OUTPUT_COLUMNS)

    # 确保输出目录存在
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 写入 Excel
    with pd.ExcelWriter(str(output_path), engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='GMV MAX汇总', index=False)

        # 获取 workbook 和 worksheet 对象进行格式化
        workbook = writer.book
        worksheet = writer.sheets['GMV MAX汇总']

        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter

        # 表头样式
        header_font = Font(bold=True, size=11, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        # 应用表头样式
        for col_idx in range(1, len(OUTPUT_COLUMNS) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # 数据行样式
        data_align = Alignment(horizontal='center', vertical='center')
        for row_idx in range(2, len(df) + 2):
            for col_idx in range(1, len(OUTPUT_COLUMNS) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.alignment = data_align
                cell.border = thin_border

        # CTR / CVR 四列：单元格为数值(百分点，如 3.25)，用数字格式显示为 '3.25%'
        pct_cols = [OUTPUT_COLUMNS.index(c) + 1 for c in
                    ('L7D CTR', 'L7D CVR', '总CTR', '总CVR')]
        for col_idx in pct_cols:
            for row_idx in range(2, len(df) + 2):
                worksheet.cell(row=row_idx, column=col_idx).number_format = '0.00"%"'

        # 自动调整列宽
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, 1):
            max_len = len(col_name)
            for row_idx in range(2, len(df) + 2):
                val = worksheet.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            # 中文字符占两个宽度，简单补偿
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = min(max_len * 1.8 + 4, 40)

        # 冻结首行
        worksheet.freeze_panes = 'A2'

    print(f"\n{'='*60}")
    print(f"输出文件: {output_path}")
    print(f"共 {len(df)} 家店铺")
    print(f"{'='*60}")
    return str(output_path)


def main():
    parser = argparse.ArgumentParser(
        description='GMV MAX 广告数据清洗汇总',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  # 指定广告数据目录
  python gmv_max_cleaner.py --ad-dir ./ad_data --mapping 人员名单.xlsx --start-date 2024-07-01

  # 指定具体文件
  python gmv_max_cleaner.py --ad-files shop001.xlsx shop002.xlsx --mapping 人员名单.xlsx

  # 不指定日期（自动检测）
  python gmv_max_cleaner.py --ad-dir ./ad_data --mapping 人员名单.xlsx
        """,
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--ad-files', nargs='+', help='广告数据文件路径列表')
    group.add_argument('--ad-dir', help='广告数据文件所在目录')
    parser.add_argument('--mapping', required=True, help='店铺人员对应名单文件路径')
    parser.add_argument('--start-date', help='L7D起始日期 (格式: 2024-07-01)')
    parser.add_argument('--output', default='GMV_MAX_汇总结果.xlsx', help='输出文件路径 (默认: GMV_MAX_汇总结果.xlsx)')
    args = parser.parse_args()

    # 收集广告数据文件
    if args.ad_dir:
        ad_dir = Path(args.ad_dir)
        if not ad_dir.exists():
            print(f"[错误] 目录不存在: {ad_dir}")
            sys.exit(1)
        ad_files = sorted(
            list(ad_dir.glob('*.xlsx')) + list(ad_dir.glob('*.xls')) + list(ad_dir.glob('*.csv'))
        )
        # 排除临时文件、人员名单文件、输出文件自身
        mapping_filename = Path(args.mapping).resolve().name
        output_filename = Path(args.output).resolve().name if hasattr(args, 'output') else ''
        ad_files = [
            f for f in ad_files
            if not f.name.startswith('~$')
            and f.resolve().name != mapping_filename
            and f.resolve().name != output_filename
        ]
        if not ad_files:
            print(f"[错误] 目录中未找到数据文件: {ad_dir}")
            sys.exit(1)
    else:
        ad_files = [Path(f) for f in args.ad_files]
        for f in ad_files:
            if not f.exists():
                print(f"[错误] 文件不存在: {f}")
                sys.exit(1)

    print(f"{'='*60}")
    print(f"GMV MAX 广告数据清洗汇总")
    print(f"{'='*60}")
    print(f"广告数据文件: {len(ad_files)} 个")
    print(f"人员名单: {args.mapping}")
    if args.start_date:
        print(f"L7D起始日期: {args.start_date}")
    else:
        print(f"L7D起始日期: 自动检测")
    print()

    # 读取人员名单
    print("--- 读取人员名单 ---")
    mapping_df = read_mapping_file(args.mapping)
    print(f"人员名单: {len(mapping_df)} 家店铺")

    mapping_codes = mapping_df['store_code'].tolist()
    mapping_dict = {}
    for _, row in mapping_df.iterrows():
        code = str(row['store_code']).strip()
        mapping_dict[code] = row.to_dict()

    print()

    # 处理每家店铺
    results = []
    processed_codes = set()

    for ad_file in ad_files:
        print(f"--- 处理: {ad_file.name} ---")

        # 提取店编
        store_code = extract_store_code(ad_file.name, mapping_codes)
        print(f"  店编: {store_code}")

        # 查找人员信息
        if store_code in mapping_dict:
            mapping_row = mapping_dict[store_code]
            print(f"  店名: {mapping_row.get('store_name', '')}")
        else:
            print(f"  [警告] 未在人员名单中找到店编 {store_code}，人员信息将为空")
            mapping_row = {
                'store_name': '', 'store_code': store_code,
                'country': '', 'junior': '', 'mid': '', 'senior': '', 'ceo': '',
            }

        processed_codes.add(store_code)

        # 读取广告数据
        try:
            ad_df = read_data_file(ad_file)
            print(f"  原始数据: {len(ad_df)} 行, {len(ad_df.columns)} 列")
        except Exception as e:
            print(f"  [错误] 读取文件失败: {e}")
            result = {col: '' for col in OUTPUT_COLUMNS}
            for key in ['L7D新建素材数', 'L7D新建素材消耗额', '总素材数', '总消耗']:
                result[key] = 0
            for key in ['L7D CTR', 'L7D CVR', '总CTR', '总CVR']:
                result[key] = 0.0
            result['店编'] = store_code
            result.update(mapping_row)
            results.append(result)
            continue

        # 智能跳过：文件不同时包含 Cost 和 Time posted 列时视为非广告数据文件
        cost_check = find_column(ad_df.columns, AD_COL_ALIASES['cost'])
        time_check = find_column(ad_df.columns, AD_COL_ALIASES['time_posted'])
        if not cost_check or not time_check:
            print(f"  [跳过] 该文件不包含广告数据所需的列(Cost/Time posted)，非广告数据文件")
            print()
            continue

        # 清洗与汇总
        start_date = args.start_date
        result = clean_and_summarize(ad_df, mapping_row, start_date)
        results.append(result)
        print()

    # 检查名单中未处理的店铺
    unprocessed = []
    for code in mapping_codes:
        if str(code).strip() not in processed_codes:
            unprocessed.append(str(code).strip())
    if unprocessed:
        print(f"[提示] 以下店铺在名单中但无数据文件: {', '.join(unprocessed)}")
        for code in unprocessed:
            row = mapping_dict.get(code, {})
            result = {col: '' for col in OUTPUT_COLUMNS}
            result['店名'] = row.get('store_name', '')
            result['店编'] = code
            result['国家'] = row.get('country', '')
            result['初级'] = row.get('junior', '')
            result['中级'] = row.get('mid', '')
            result['储高/见高'] = row.get('senior', '')
            result['CEO'] = row.get('ceo', '')
            for key in ['L7D新建素材数', 'L7D新建素材消耗额', '总素材数', '总消耗']:
                result[key] = 0
            for key in ['L7D CTR', 'L7D CVR', '总CTR', '总CVR']:
                result[key] = 0.0
            results.append(result)
        print()

    # 输出汇总表
    output_path = write_output(results, args.output)

    # 打印汇总预览
    print("\n汇总预览:")
    df = pd.DataFrame(results, columns=OUTPUT_COLUMNS)
    pd.set_option('display.max_columns', None)
    pd.set_option('display.width', 200)
    pd.set_option('display.max_colwidth', 15)
    print(df.to_string(index=False))


if __name__ == '__main__':
    main()
