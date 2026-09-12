"""TikTok native snapshots. Separate from the immutable original Mabang pipeline.

Mappings/keys/date formats must come from a reviewed export, never guessed aliases.
Only normalized allowlisted fields leave this module for an LLM.
"""
import csv
import hashlib
import io
import json
from collections import Counter
from datetime import datetime, date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook

FIELDS = {
    'orders': {'order_id', 'line_id', 'sku', 'product', 'status', 'created_at', 'quantity', 'currency', 'returned_quantity', 'return_type'},
    'settlement': {'transaction_id', 'order_id', 'related_order_id', 'settled_at', 'transaction_type', 'settlement', 'revenue', 'refund', 'currency', 'fees', 'adjustment', 'commission', 'transaction_fee', 'shipping_fee', 'affiliate_fee'},
}
NEEDED = {'orders': {'order_id', 'created_at', 'status'},
          'settlement': {'transaction_id', 'settled_at', 'settlement'}}
NUMBERS = {'quantity', 'returned_quantity', 'settlement', 'revenue', 'refund', 'fees', 'adjustment', 'commission', 'transaction_fee', 'shipping_fee', 'affiliate_fee'}


def fail(message):
    raise HTTPException(422, message)


def report_kinds(record):
    """Legacy jobs retain their original scope. New jobs explicitly request a bill only."""
    kinds=record.get('required_reports',['orders','settlement'])
    if kinds not in (['settlement'],['orders','settlement']):fail('采集报表范围无效')
    return kinds


def parse_report(content, filename, kind, spec, start, end, currency):
    """Parse strict, reviewed column positions. Raw bytes are saved by the caller first."""
    if kind not in FIELDS: fail('未知原生报表类型')
    if not spec.get('headers') or not spec.get('unique_key'): fail('尚未验证原始表头与唯一键')
    fields = spec.get('fields', {})
    if not NEEDED[kind] <= fields.keys() or not set(fields) <= FIELDS[kind]:
        fail('字段映射缺少必要字段或包含非白名单字段')
    if not set(spec['unique_key']) <= fields.keys(): fail('唯一键必须使用已映射字段')
    suffix = Path(filename).suffix.lower()
    control = {}
    control_labels=spec.get('control_labels',{'period':'时间范围','timezone':'时区','currency':'货币'})
    try:
        if suffix == '.csv':
            matrix = list(csv.reader(io.StringIO(content.decode('utf-8-sig'))))
            sheet = 'CSV'
        elif suffix == '.xlsx':
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = spec.get('sheet', '')
                if sheet not in book.sheetnames: fail('工作表与已验证规范不同')
                book[sheet].reset_dimensions()
                matrix = [[c.value if c.value is not None else '' for c in row] for row in book[sheet].rows]
                if spec.get('control_sheet'):
                    if spec['control_sheet'] not in book.sheetnames: fail('缺少平台汇总对账页')
                    ws = book[spec['control_sheet']]; ws.reset_dimensions()
                    labels = {*control_labels.values(),*spec.get('control_totals',{}).values()}
                    for r in ws.values:
                        cells = [str(v).strip() for v in r if v is not None and str(v).strip()]
                        if cells and cells[0] in labels:
                            if len(cells)!=2 or cells[0] in control: fail('汇总控制字段重复或格式改变')
                            control[cells[0]] = cells[1]
            finally:
                book.close()
        else: fail('原生采集只接受 XLSX 或 UTF-8 CSV，不能接受网页或临时下载文件')
    except HTTPException: raise
    except Exception: fail('报表无法完整读取，请检查下载或编码')
    offset = spec.get('header_row', 1) - 1
    if offset < 0 or offset >= len(matrix): fail('表头位置无效')
    header = [str(v).strip() for v in matrix[offset]]
    if header != spec['headers'] or len(set(h for h in header if h)) != len([h for h in header if h]):
        fail('表头顺序发生变化，需要重新审核；未覆盖当前数据')
    if not set(fields.values()) <= set(header) or '' in fields.values(): fail('映射字段不在报表内或映射到空白分隔列')
    indexes = {k: header.index(v) for k, v in fields.items()}
    result, keys = [], set()
    fee_fields=spec.get('fee_fields',[])
    if not set(fee_fields)<=set(header) or '' in fee_fields:fail('费用拆解字段不在已审核报表中')
    fee_values={name:[] for name in fee_fields}
    start, end = date.fromisoformat(start), date.fromisoformat(end)
    for source_row, raw in enumerate(matrix[offset+1:], offset+2):
        if not any(v not in ('', None) for v in raw): continue
        if len(raw) != len(header): fail(f'第 {source_row} 行列数与表头不同')
        for name in fee_fields:
            value=raw[header.index(name)]
            try:
                number=None if value in ('',None) else Decimal(str(value).strip())
                if number is not None and not number.is_finite():raise InvalidOperation()
            except InvalidOperation:fail('费用字段不是有效数值：'+name)
            fee_values[name].append(number)
        item = {'source_row': source_row}
        for key, col in indexes.items():
            value = raw[col]
            if key.endswith('_id') or key == 'sku':
                if isinstance(value, (int, float)) and abs(value) >= 10**15:
                    fail(f'第 {source_row} 行长 ID 以数字保存，可能已损失精度')
                if isinstance(value, float):
                    if not value.is_integer(): fail('ID 不应为小数')
                    value = int(value)
            if key in NUMBERS:
                if value in ('', None): item[key] = None; continue
                try:
                    n = Decimal(str(value).replace(',', ''))
                    if not n.is_finite(): raise InvalidOperation()
                    if key == 'quantity' and (n < 0 or n != n.to_integral_value()): raise InvalidOperation()
                except (InvalidOperation, ValueError): fail(f'第 {source_row} 行 {key} 不是有效数值')
                item[key] = str(n)
            elif key in ('created_at', 'settled_at'):
                try:
                    dt = value if isinstance(value, datetime) else datetime.strptime(str(value).strip(), spec['date_format'])
                    if dt.tzinfo is not None: fail('带时区日期需先验证平台时区转换，不自动截断')
                    if not start <= dt.date() <= end: fail(f'第 {source_row} 行日期不在本次筛选范围')
                    item[key] = dt.isoformat()
                except (ValueError, KeyError): fail(f'第 {source_row} 行日期格式未通过验证')
            else:
                item[key] = str(value or '').strip()
        if kind=='settlement' and spec.get('order_link_by_type'):
            item['order_id'] = item['transaction_id'] if item.get('transaction_type')==spec['order_link_by_type'] else item.get('related_order_id','')
        key = tuple(item.get(k) for k in spec['unique_key'])
        if any(v in ('', None) for v in key): fail(f'第 {source_row} 行缺少唯一键')
        if key in keys: fail('唯一键重复，请核实报表粒度；不自动去掉可能不同的交易行')
        keys.add(key)
        if item.get('currency', currency) != currency: fail('报表币种与绑定币种不同')
        result.append(item)
    if not result and not spec.get('empty_export_verified'): fail('空报表需要验证平台确实为零，不能将失败下载视为零数据')
    if spec.get('control_sheet'):
        if control.get(control_labels['period']) != start.strftime('%Y/%m/%d')+'-'+end.strftime('%Y/%m/%d') or control.get(control_labels['timezone'])!=spec.get('control_timezone','UTC+8') or control.get(control_labels['currency'])!=currency:
            fail('平台汇总页的周期、时区或币种与采集任务不同')
        for key,label in spec['control_totals'].items():
            amount = total(result,key,fields)
            try:
                if amount is None or Decimal(amount)!=Decimal(control[label]): fail('结算明细与平台汇总金额不一致：'+label)
            except (KeyError,InvalidOperation): fail('缺少有效平台汇总金额：'+label)
    fee_breakdown={name:None if any(v is None for v in values) else str(sum(values,Decimal(0))) for name,values in fee_values.items()}
    return {'kind': kind, 'header': header, 'sheet': sheet, 'row_count': len(result), 'rows': result, 'control':control,'fee_breakdown':fee_breakdown,
            'sha256': hashlib.sha256(content).hexdigest(), 'name': filename,
            'date_basis': '创建日期' if kind == 'orders' else '结算日期', 'mapped_fields': sorted(fields)}


def total(rows, key, available):
    if key not in available or any(r.get(key) is None for r in rows): return None
    return str(sum((Decimal(r[key]) for r in rows), Decimal(0)))


def metrics(snapshot):
    has_orders='orders' in snapshot['reports']
    orders, settled = snapshot['reports'].get('orders',{'rows':[],'mapped_fields':[]}), snapshot['reports']['settlement']
    rows, tx = orders['rows'], settled['rows']
    ids = {r['order_id'] for r in rows}
    outside = [r for r in tx if not r.get('order_id') or r['order_id'] not in ids]
    breakdown=operating_breakdown(snapshot)
    return {
        'shop': snapshot['shop'], 'period': snapshot['start']+'/'+snapshot['end'],
        'currency': snapshot['currency'], 'platform_timezone': snapshot['platform_timezone'],
        'orders_available':has_orders, 'report_name':'账单',
        'partial_day':snapshot.get('partial_day'),
        'completeness_note':'包含当天数据；当天仅截至导出时，不是完整自然日。' if snapshot.get('partial_day') else '按原批次日期范围；不补造缺失数据。',
        'orders': {'order_count': len(ids) if has_orders else None, 'line_count': len(rows) if has_orders else None,
                   'quantity': total(rows, 'quantity', orders['mapped_fields']),
                   'returned_quantity': total(rows,'returned_quantity',orders['mapped_fields']),
                   'line_status_counts': dict(Counter(r['status'] for r in rows))},
        'daily':breakdown['daily'], 'product_risk_count':len(breakdown['anomalies']) if has_orders else None,
        'product_risk_preview':breakdown['anomalies'][:5],
        'settlement': {'transaction_count': len(tx),
                       **{k: total(tx, k, settled['mapped_fields']) for k in ('settlement', 'revenue', 'refund', 'fees', 'adjustment', 'commission', 'transaction_fee', 'shipping_fee', 'affiliate_fee')},
                       'transaction_types':dict(Counter(r.get('transaction_type','未提供') for r in tx)),
                       'fee_breakdown':settled.get('fee_breakdown',{}),
                       'unmatched_transaction_count': len(outside) if has_orders else None},
        'rules': ['新增 TikTok 原生口径，不替代马帮订单或原利润计算。',
                  '本批次只包含财务交易账单，不含订单列表；不能从账单推算总订单量、日销、取消率。' if not has_orders else '本历史批次同时包含订单列表与账单。',
                  '订单按创建日期，结算按结算日期；两个窗口不是同一订单群。',
                  '状态计数为商品行数，不是订单状态数；结算金额保留原符号，不称利润。',
                  '未匹配交易仅表示当前创建日期窗口内未匹配；不代表错账。',
                  '缺少费用、退款或数量字段时返回 null；不估算退款率或原因。',
                  '费用拆解按平台原列独立展示，含父子层级项目，不能把这些列再次加总作为总费用。'],
    }


def operating_breakdown(snapshot):
    """New, named screening rules. Counts use native line scope, not legacy daily sales."""
    if 'orders' not in snapshot['reports']:return {'products':[],'anomalies':[],'daily':[]}
    orders=snapshot['reports']['orders']; groups={};daily={}
    start=date.fromisoformat(snapshot['start']);end=date.fromisoformat(snapshot['end'])
    while start<=end:
        daily[start.isoformat()]={'date':start.isoformat(),'ids':set(),'quantity':Decimal(0),'missing_quantity':False}
        start+=timedelta(days=1)
    for r in orders['rows']:
        day=daily[r['created_at'][:10]];day['ids'].add(r['order_id'])
        day['missing_quantity']|=r.get('quantity') is None
        day['quantity']+=Decimal(r.get('quantity') or 0)
        key=r.get('line_id') or r.get('sku') or '未提供 SKU'
        g=groups.setdefault(key,{'sku_id':key,'sku':r.get('sku',''),'product':r.get('product',''),'line_count':0,'cancelled_lines':0,'returned_quantity':Decimal(0),'source_rows':[]})
        g['line_count']+=1;g['cancelled_lines']+=r['status'] in ('已取消','Cancelled','Canceled')
        g['returned_quantity']+=Decimal(r.get('returned_quantity') or 0);g['source_rows'].append(r['source_row'])
    products=[]
    for g in groups.values():
        g['returned_quantity']=str(g['returned_quantity']) if 'returned_quantity' in orders['mapped_fields'] else None
        g['cancelled_line_share']=str(Decimal(g['cancelled_lines'])/g['line_count'])
        g['rule']='新增筛查：取消商品行≥5，且取消商品行/创建窗口商品行≥20%；不是退款率或原因诊断'
        g['risk']=g['cancelled_lines']>=5 and Decimal(g['cancelled_line_share'])>=Decimal('.2')
        products.append(g)
    products.sort(key=lambda g:(-g['cancelled_lines'],g['sku_id']))
    # Full source rows stay in paginated tools; automatic summary needs only representative references.
    anomalies=[{**g,'source_rows':g['source_rows'][:3]} for g in products if g['risk']]
    return {'products':products,'anomalies':anomalies,'daily':[{'date':d['date'],'order_count':len(d['ids']),'quantity':None if d['missing_quantity'] else str(d['quantity'])} for d in daily.values()]}


def native_query(current, store_id='', kind='summary', offset=0, limit=50):
    snapshots = current['data'].get('tiktok_native', {})
    if not store_id:
        if len(snapshots) != 1:
            return {'version': current['version'], 'stores': list(snapshots),
                    'store_options':[{'store_id':sid,'shop':s['shop'],'currency':s['currency'],'period':s['start']+'/'+s['end']} for sid,s in snapshots.items()],
                    'note': '请选择精确店铺环境 ID；多店币种和批次分别查询，不合计'}
        store_id = next(iter(snapshots))
    snap = snapshots.get(store_id)
    if not snap: return {'version': current['version'], 'total': 0, 'note': '当前版本没有该店铺的原生采集数据'}
    out = {'version': current['version'], 'job_id': snap['job_id'], 'summary': metrics(snap)}
    if kind in ('orders','products','anomalies') and 'orders' not in snap['reports']:
        return {**out,'kind':kind,'available':False,'total':None,'rows':[],'note':'本批次只导出账单，未采集订单列表，不能计算订单或 SKU 销售指标。'}
    if kind in ('orders', 'settlement'):
        report = snap['reports'][kind]
        out.update(kind=kind, total=len(report['rows']), offset=offset,
                   rows=report['rows'][offset:offset+limit], file_sha256=report['sha256'], sheet=report['sheet'])
    elif kind in ('products','anomalies'):
        rows=operating_breakdown(snap)[kind]
        out.update(kind=kind,total=len(rows),offset=offset,rows=rows[offset:offset+limit],file_sha256=snap['reports']['orders']['sha256'],sheet=snap['reports']['orders']['sheet'])
    return out


def content_fingerprint(reports):
    # Excel container bytes can change without any business values changing.
    value = {k: {**{f:v[f] for f in ('header','sheet','mapped_fields')},
                'rows':sorted([{key:value for key,value in row.items() if key!='source_row'} for row in v['rows']],key=lambda r:json.dumps(r,sort_keys=True))} for k,v in reports.items()}
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
