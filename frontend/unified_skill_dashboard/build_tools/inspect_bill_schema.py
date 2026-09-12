"""Read-only original XLSX schema/control-sheet inspection for onboarding."""
import sys,json
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    import server
    c=server.collection;b=c.batch(sys.argv[1])
    for rid in b['run_ids']:
        j=c.get(rid);f=j.get('files',{}).get('settlement')
        if not f:continue
        wb=load_workbook(c.hub.data/'raw/ai_workbench/imports'/rid/f['saved_name'],read_only=True,data_only=False)
        out={'id':rid,'store_id':j['store_id'],'shop':j['shop'],'file':f['name'],'sheets':wb.sheetnames,'tables':{}}
        for ws in wb:
            ws.reset_dimensions();rows=ws.values
            if ws.title in ('Report','Reports','报告'):
                out['tables'][ws.title]=[[v for v in r if v not in (None,'')] for r in rows if any(v not in (None,'') for v in r)][:30]
            else:
                first=next(rows,());sample=next(rows,())
                out['tables'][ws.title]={'headers':first,'row_count':1+sum(1 for r in rows if any(v is not None for v in r))}
                if ws.title in ('Order details','订单详情'):out['tables'][ws.title]['type_date_currency']=sample[1:5]
        wb.close();print(json.dumps(out,ensure_ascii=False,default=str),flush=True)
